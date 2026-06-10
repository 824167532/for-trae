"""
客户路由
"""
from flask import Blueprint, request, jsonify
from io import BytesIO
from openpyxl import load_workbook
from services.customer_service import CustomerService

customers_bp = Blueprint('customers', __name__, url_prefix='/api/customers')
customer_service = CustomerService()


@customers_bp.route('', methods=['GET'])
def get_customers():
    """获取所有客户"""
    customers = customer_service.get_all()
    return jsonify(customers)


@customers_bp.route('', methods=['POST'])
def create_customer():
    """创建客户"""
    data = request.get_json()
    display_name = data.get('display_name')
    relative_path = data.get('relative_path')
    
    if not display_name or not relative_path:
        return jsonify({'error': '缺少必要参数'}), 400
    
    customer = customer_service.create(display_name, relative_path)
    if customer:
        return jsonify(customer)
    else:
        return jsonify({'error': '客户名称已存在'}), 400


@customers_bp.route('/<int:customer_id>', methods=['PUT'])
def update_customer(customer_id):
    """更新客户"""
    data = request.get_json()
    display_name = data.get('display_name')
    relative_path = data.get('relative_path')
    
    if not display_name or not relative_path:
        return jsonify({'error': '缺少必要参数'}), 400
    
    customer = customer_service.update(customer_id, display_name, relative_path)
    if customer:
        return jsonify(customer)
    else:
        return jsonify({'error': '客户名称已存在或客户不存在'}), 400


@customers_bp.route('/<int:customer_id>', methods=['DELETE'])
def delete_customer(customer_id):
    """删除客户"""
    success = customer_service.delete(customer_id)
    if success:
        return jsonify({'success': True})
    else:
        return jsonify({'error': '删除失败'}), 400


@customers_bp.route('/import', methods=['POST'])
def import_customers():
    """导入客户"""
    if 'file' not in request.files:
        return jsonify({'error': '未上传文件'}), 400
    
    file = request.files['file']
    if not file.filename.endswith('.xlsx'):
        return jsonify({'error': '仅支持xlsx格式'}), 400
    
    try:
        # 读取Excel文件
        wb = load_workbook(BytesIO(file.read()))
        ws = wb.active
        
        # 解析数据
        data_list = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                data_list.append({
                    '客户显示名': str(row[0]),
                    '相对路径': str(row[1])
                })
        
        # 导入数据
        result = customer_service.import_from_excel(data_list)
        return jsonify(result)
    
    except Exception as e:
        return jsonify({'error': str(e)}), 400


@customers_bp.route('/template', methods=['GET'])
def export_template():
    """导出客户导入模板"""
    from openpyxl import Workbook
    from flask import send_file
    
    wb = Workbook()
    ws = wb.active
    ws.title = '客户列表'
    
    # 设置表头
    ws['A1'] = '客户显示名'
    ws['B1'] = '相对路径'
    
    # 设置示例数据
    ws['A2'] = '广东移动'
    ws['B2'] = '运营商\\移动\\广东'
    ws['A3'] = '广东移动（深圳）'
    ws['B3'] = '运营商\\移动\\广东\\深圳'
    
    # 保存到内存
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='客户导入模板.xlsx'
    )